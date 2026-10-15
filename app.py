import io
import os
from datetime import datetime, timedelta, timezone

import openpyxl
import pandas as pd
import streamlit as st
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Calculadora de Tarimas", layout="centered")

st.title("Calculadora de Porcentajes por Tarima")

# Archivo de almacenamiento persistente
EXCEL_FILE = "registros_tarimas.xlsx"

# Guatemala está en UTC-6 todo el año (no usa horario de verano).
# Así la hora sale bien aunque el servidor esté en otra zona horaria.
ZONA_GT = timezone(timedelta(hours=-6))

DIAS_SEMANA = {
    "Monday": "Lunes",
    "Tuesday": "Martes",
    "Wednesday": "Miércoles",
    "Thursday": "Jueves",
    "Friday": "Viernes",
    "Saturday": "Sábado",
    "Sunday": "Domingo",
}

COLUMNAS = [
    "Día",
    "Fecha",
    "Hora",
    "Nº / ID Tarima",
    "Peso Total (FI)",
    "Merma",
    "% Merma",
    "Bueno",
    "% Bueno",
    "No Colocable",
    "% No Colocable",
]

PARES_PESO_PCT = [
    ("Merma", "% Merma"),
    ("Bueno", "% Bueno"),
    ("No Colocable", "% No Colocable"),
]


def a_numero(valor):
    """Convierte a float; si está vacío/NaN devuelve 0."""
    try:
        if pd.isna(valor):
            return 0.0
        return float(valor)
    except (TypeError, ValueError):
        return 0.0


def cargar_datos():
    """Carga el Excel y arregla de una vez los registros ya guardados:
    - Renombra columnas viejas ("Procesado" -> "No Colocable")
    - Rellena datos faltantes y recalcula los % que falten
    - Guarda el archivo corregido automáticamente
    """
    if not os.path.exists(EXCEL_FILE):
        return pd.DataFrame(columns=COLUMNAS)

    df = pd.read_excel(EXCEL_FILE)
    cambio = False

    # Migrar columnas viejas
    for viejo, nuevo in [
        ("Procesado", "No Colocable"),
        ("% Procesado", "% No Colocable"),
    ]:
        if viejo in df.columns:
            if nuevo in df.columns:
                df[nuevo] = df[nuevo].combine_first(df[viejo])
            else:
                df[nuevo] = df[viejo]
            df = df.drop(columns=[viejo])
            cambio = True

    # Asegurar todas las columnas
    for col in COLUMNAS:
        if col not in df.columns:
            df[col] = None
            cambio = True

    df = df[COLUMNAS].copy()

    # Rellenar pesos vacíos con 0 y recalcular % faltantes
    for idx, fila in df.iterrows():
        fi = a_numero(fila["Peso Total (FI)"])
        for peso_col, pct_col in PARES_PESO_PCT:
            if pd.isna(fila[peso_col]):
                df.at[idx, peso_col] = 0.0
                cambio = True
            if pd.isna(fila[pct_col]):
                peso = a_numero(df.at[idx, peso_col])
                pct = (peso / fi * 100) if fi > 0 else 0.0
                df.at[idx, pct_col] = f"{pct:.1f}%"
                cambio = True

    if cambio:
        df.to_excel(EXCEL_FILE, index=False)

    return df


def generar_excel_estilizado(df):
    output = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reporte de Tarimas"

    ws.views.sheetView[0].showGridLines = True

    # 1. Título del Reporte
    ws.merge_cells("A1:K1")
    title_cell = ws["A1"]
    title_cell.value = "REPORTE Y CONTROL DE TARIMAS"
    title_cell.font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(
        start_color="1F4E79", end_color="1F4E79", fill_type="solid"
    )
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 35

    # 2. Encabezados con colores por categoría
    headers = [
        "Nº / ID Tarima",
        "Día",
        "Fecha",
        "Hora",
        "Peso Total (FI)",
        "Merma",
        "% Merma",
        "Bueno",
        "% Bueno",
        "No Colocable",
        "% No Colocable",
    ]

    colores = [
        "2F5597",  # Azul - Identificación
        "2F5597",
        "2F5597",
        "2F5597",
        "215E21",  # Verde oscuro - Peso Total
        "C00000",  # Rojo - Merma
        "C00000",
        "385723",  # Verde - Bueno
        "385723",
        "C55A11",  # Naranja - No Colocable
        "C55A11",
    ]
    header_fills = [
        PatternFill(start_color=c, end_color=c, fill_type="solid")
        for c in colores
    ]

    ws.row_dimensions[3].height = 28

    header_border = Border(
        left=Side(style="thin", color="FFFFFF"),
        right=Side(style="thin", color="FFFFFF"),
        top=Side(style="medium", color="1F4E79"),
        bottom=Side(style="medium", color="1F4E79"),
    )

    for col_num, (header_text, fill) in enumerate(
        zip(headers, header_fills), 1
    ):
        cell = ws.cell(row=3, column=col_num, value=header_text)
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        cell.border = header_border

    # 3. Filas de datos
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    zebra_fill = PatternFill(
        start_color="F2F2F2", end_color="F2F2F2", fill_type="solid"
    )

    def clean_pct(val):
        if isinstance(val, str):
            try:
                return float(val.replace("%", "").strip()) / 100.0
            except ValueError:
                return 0.0
        val_float = a_numero(val)
        return val_float / 100.0 if val_float > 1 else val_float

    def texto(val):
        return "" if pd.isna(val) else val

    for i, (_, row) in enumerate(df.iterrows()):
        current_row = i + 4
        ws.row_dimensions[current_row].height = 22
        is_even = current_row % 2 == 0

        valores = [
            texto(row.get("Nº / ID Tarima", "")),
            texto(row.get("Día", "")),
            texto(row.get("Fecha", "")),
            texto(row.get("Hora", "")),
            a_numero(row.get("Peso Total (FI)", 0)),
            a_numero(row.get("Merma", 0)),
            clean_pct(row.get("% Merma", 0)),
            a_numero(row.get("Bueno", 0)),
            clean_pct(row.get("% Bueno", 0)),
            a_numero(row.get("No Colocable", 0)),
            clean_pct(row.get("% No Colocable", 0)),
        ]

        for col_idx, val in enumerate(valores, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = Font(name="Calibri", size=11)
            cell.border = thin_border

            if is_even:
                cell.fill = zebra_fill

            if col_idx in [1, 2, 3, 4]:
                cell.alignment = Alignment(
                    horizontal="center", vertical="center"
                )
            elif col_idx in [5, 6, 8, 10]:
                cell.alignment = Alignment(
                    horizontal="right", vertical="center"
                )
                cell.number_format = "#,##0.0"
            elif col_idx in [7, 9, 11]:
                cell.alignment = Alignment(
                    horizontal="right", vertical="center"
                )
                cell.number_format = "0.0%"

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = 16

    wb.save(output)
    return output.getvalue()


# --- ENTRADAS DE DATOS ---
st.subheader("Ingreso de Datos de la Tarima")

ahora_vista = datetime.now(ZONA_GT)
st.caption(
    f"🕒 {DIAS_SEMANA[ahora_vista.strftime('%A')]} "
    f"{ahora_vista.strftime('%Y-%m-%d')} · Hora actual: "
    f"{ahora_vista.strftime('%H:%M:%S')} (se registra al guardar)"
)

tarima_id = st.text_input("Número / ID de Tarima o Lote:", value="")
fi = st.number_input(
    "Peso Total Inicial (FI):", min_value=0.0, value=0.0, step=0.1
)
merma = st.number_input("Peso de Merma:", min_value=0.0, value=0.0, step=0.1)
bueno = st.number_input("Peso de Bueno:", min_value=0.0, value=0.0, step=0.1)
no_colocable = st.number_input(
    "Peso No Colocable:", min_value=0.0, value=0.0, step=0.1
)

# VALIDACIONES
if fi <= 0:
    st.info(
        "Ingresa un Peso Total Inicial (FI) mayor a 0 para calcular los porcentajes."
    )
else:
    suma_partes = merma + bueno + no_colocable
    restante = fi - suma_partes  # se va restando en vivo

    # Resta en vivo
    c1, c2, c3 = st.columns(3)
    c1.metric("Peso Total (FI)", f"{fi:,.1f}")
    c2.metric("Ya asignado", f"{suma_partes:,.1f}")
    c3.metric("Restante por asignar", f"{restante:,.1f}")

    if restante < 0:
        st.error(
            f"Error: La suma de las partes ({suma_partes:.2f}) sobrepasa el "
            f"Peso Total Inicial ({fi:.2f}). Te pasaste por {abs(restante):.2f}."
        )
    else:
        pct_merma = (merma / fi) * 100
        pct_bueno = (bueno / fi) * 100
        pct_no_colocable = (no_colocable / fi) * 100
        pct_restante = (restante / fi) * 100

        data_preview = {
            "Concepto": ["Merma", "Bueno", "No Colocable", "Sin asignar"],
            "Peso (lb/kg)": [merma, bueno, no_colocable, restante],
            "Porcentaje (%)": [
                f"{pct_merma:.1f}%",
                f"{pct_bueno:.1f}%",
                f"{pct_no_colocable:.1f}%",
                f"{pct_restante:.1f}%",
            ],
        }
        st.subheader("Resumen de Porcentajes")
        st.table(pd.DataFrame(data_preview))

        if st.button("Guardar Registro en Excel"):
            df_existente = cargar_datos()

            ahora = datetime.now(ZONA_GT)
            dia_esp = DIAS_SEMANA[ahora.strftime("%A")]
            fecha_str = ahora.strftime("%Y-%m-%d")
            hora_str = ahora.strftime("%H:%M:%S")

            nuevo_registro = pd.DataFrame(
                [
                    {
                        "Día": dia_esp,
                        "Fecha": fecha_str,
                        "Hora": hora_str,
                        "Nº / ID Tarima": (
                            tarima_id if tarima_id else "Sin Identificador"
                        ),
                        "Peso Total (FI)": fi,
                        "Merma": merma,
                        "% Merma": f"{pct_merma:.1f}%",
                        "Bueno": bueno,
                        "% Bueno": f"{pct_bueno:.1f}%",
                        "No Colocable": no_colocable,
                        "% No Colocable": f"{pct_no_colocable:.1f}%",
                    }
                ]
            )

            df_actualizado = pd.concat(
                [df_existente, nuevo_registro], ignore_index=True
            )[COLUMNAS]
            df_actualizado.to_excel(EXCEL_FILE, index=False)
            st.success("¡Registro guardado exitosamente!")
            st.rerun()

# --- HISTORIAL Y BORRADO ---
st.divider()
st.subheader("Historial de Registros Guardados")

df_historial = cargar_datos()  # también corrige y guarda los registros viejos

if not df_historial.empty:
    st.dataframe(df_historial)

    st.markdown("### Eliminar un registro")

    opciones_eliminar = [
        f"Fila {idx + 1} - {row['Nº / ID Tarima']} ({row['Fecha']} {row['Hora']})"
        for idx, row in df_historial.iterrows()
    ]

    registro_a_eliminar = st.selectbox(
        "Selecciona la fila a eliminar:", opciones_eliminar
    )

    if st.button("Eliminar Fila Seleccionada"):
        indice_eliminar = opciones_eliminar.index(registro_a_eliminar)
        df_historial = df_historial.drop(index=indice_eliminar).reset_index(
            drop=True
        )
        df_historial.to_excel(EXCEL_FILE, index=False)
        st.success("¡Registro eliminado!")
        st.rerun()

    st.divider()

    excel_binario = generar_excel_estilizado(df_historial)

    st.download_button(
        label="Descargar Reporte Formateado en Excel",
        data=excel_binario,
        file_name="Reporte_Tarimas_Formateado.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
else:
    st.info("Aún no hay registros guardados en el historial.")
